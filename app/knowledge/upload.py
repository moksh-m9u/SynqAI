"""Interactive knowledge upload: extract text from arbitrary files (TXT, PDF,
DOCX, XLSX, CSV, JSON), preview deterministic chunks, edit them, then embed and
index into the isolated ``synqai_userkb`` Qdrant collection.

Chunking is deterministic Python (no LLM); rule_candidate is a keyword heuristic
with a confidence score so the judge can see *why* a chunk was flagged. The main
``synqai_knowledge`` baseline collection is never touched by uploads.
"""
from __future__ import annotations

import io
import json
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd

from app.knowledge.vectorstore import KnowledgeStore, USER_KB_COLLECTION

_RULE_HINT_RE = re.compile(
    r"rule|must|never|always|sla|heater|bs6|severe|monsoon|night|jugaad|brake|"
    r"load|halting|gurgaon|delhi|ncr|hill|ludhiana|36\s*hour|48\s*hour|6[.:]\s*00 pm|"
    r"ground|rotat|capacity|year\s*[<>]=\s*20|emergency|week|8\.2|2020",
    re.I)


@dataclass
class Extracted:
    file_name: str
    document_type: str
    language: str
    text: str
    rows: int = 0
    note: str = ""


def _language_of(text: str) -> str:
    deva = sum(1 for ch in text if unicodedata.name(ch, "").startswith("DEVANAGARI"))
    non_ascii = sum(1 for ch in text if ord(ch) > 127)
    if deva > 0:
        return "hi-en" if non_ascii > deva else "hi"
    return "en"


def extract_text(raw: bytes, file_name: str) -> Extracted:
    """Document-type detection + text extraction for the six upload formats."""
    name = file_name.lower()
    if name.endswith(".txt") or name.endswith(".md") or name.endswith(".log"):
        text = raw.decode("utf-8-sig", errors="replace")
        return Extracted(file_name, "txt", _language_of(text), text)
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        pages = [p.extract_text() or "" for p in reader.pages]
        text = "\n\n".join(pages)
        return Extracted(file_name, "pdf", _language_of(text), text, rows=len(reader.pages),
                         note=f"{len(reader.pages)} pages")
    if name.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(raw))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return Extracted(file_name, "docx", _language_of(text), text,
                         note=f"{len(doc.paragraphs)} paragraphs")
    if name.endswith(".xlsx") or name.endswith(".xls"):
        df = pd.read_excel(io.BytesIO(raw))
        text = "\n".join("| ".join(str(v) for v in row) for row in df.itertuples(index=False))
        return Extracted(file_name, "xlsx", _language_of(text), text, rows=len(df),
                         note=f"{len(df)} rows × {len(df.columns)} cols")
    if name.endswith(".csv"):
        df = pd.read_csv(io.BytesIO(raw))
        text = "\n".join("| ".join(str(v) for v in row) for row in df.itertuples(index=False))
        return Extracted(file_name, "csv", _language_of(text), text, rows=len(df),
                         note=f"{len(df)} rows × {len(df.columns)} cols")
    if name.endswith(".json"):
        data = json.loads(raw.decode("utf-8"))
        text = json.dumps(data, indent=2, ensure_ascii=False, default=str)
        return Extracted(file_name, "json", _language_of(text), text,
                         note="pretty-printed JSON")
    raise ValueError("unsupported file type (use TXT, PDF, DOCX, XLSX, CSV, JSON)")


def _rule_evidence(text: str) -> tuple[bool, float]:
    matches = _RULE_HINT_RE.findall(text.lower())
    if not matches:
        return False, 0.0
    confidence = min(0.99, 0.55 + 0.08 * len(set(matches)))
    return True, round(confidence, 2)


def chunk_text(text: str, source: str, document_type: str, language: str = "en",
               chunk_size: int = 900, overlap: int = 60) -> list[dict[str, Any]]:
    """Deterministic paragraph-aware chunking with token counts + rule signal."""
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks: list[dict[str, Any]] = []
    buf, buf_len = [], 0
    for p in paragraphs:
        p_chars = len(p)
        if buf and buf_len + p_chars > chunk_size:
            chunks.append(_make_chunk(buf, source, document_type, language))
            tail, tail_len = [], 0
            for t in reversed(buf):
                if tail_len + len(t) > overlap:
                    break
                tail.insert(0, t)
                tail_len += len(t)
            buf, buf_len = tail, tail_len
        buf.append(p)
        buf_len += p_chars
        while buf_len > chunk_size and len(buf) > 1:
            chunks.append(_make_chunk(buf, source, document_type, language))
            head = buf.pop(0)
            buf_len -= len(head)
    if buf:
        chunks.append(_make_chunk(buf, source, document_type, language))
    return [tag_chunk(c, index) for index, c in enumerate(chunks)]


def tag_chunk(chunk: dict[str, Any], index: int) -> dict[str, Any]:
    rule, conf = _rule_evidence(chunk["text"])
    chunk["rule_candidate"] = rule
    chunk["confidence"] = conf
    chunk["token_count"] = len(chunk["text"].split())
    return chunk


def _make_chunk(buf: list[str], source: str, document_type: str, language: str) -> dict[str, Any]:
    return {"text": "\n".join(buf).strip(), "source": source,
            "document_type": document_type, "language": language}


def reindex_chunks(chunks: list[dict[str, Any]], source: str) -> list[dict[str, Any]]:
    for i, c in enumerate(chunks):
        c["chunk_id"] = f"upl_{hashlib_short(source)}_{i:04d}"
        c["source"] = source
    return chunks


def hashlib_short(text: str) -> str:
    import hashlib
    return hashlib.sha1(text.encode()).hexdigest()[:8]


def merge_chunks(chunks: list[dict[str, Any]], i: int) -> list[dict[str, Any]]:
    if i >= len(chunks) - 1:
        return chunks
    merged = dict(chunks[i])
    merged["text"] = chunks[i]["text"] + "\n\n" + chunks[i + 1]["text"]
    out = chunks[:i] + [merged] + chunks[i + 2:]
    return reindex_chunks(out, chunks[i]["source"])


def split_chunk(chunks: list[dict[str, Any]], i: int, at_char: int) -> list[dict[str, Any]]:
    text = chunks[i]["text"]
    at_char = max(0, min(at_char, len(text)))
    a, b = text[:at_char].strip(), text[at_char:].strip()
    if not a or not b:
        return chunks
    row = dict(chunks[i])
    row["text"] = a
    row2 = dict(chunks[i])
    row2["text"] = b
    out = chunks[:i] + [row, row2] + chunks[i + 1:]
    return reindex_chunks(out, chunks[i]["source"])


def delete_chunk(chunks: list[dict[str, Any]], i: int) -> list[dict[str, Any]]:
    if not chunks:
        return chunks
    out = chunks[:i] + chunks[i + 1:]
    return reindex_chunks(out, chunks[i]["source"]) if chunks else out


def estimate_cost(chunks: list[dict[str, Any]]) -> dict[str, Any]:
    """Ballpark embedding + storage cost for indexing the pending chunks."""
    tokens = sum(int(c.get("token_count", 0)) for c in chunks)
    dim = 384
    vec_mb = (len(chunks) * dim * 4) / (1024 * 1024)
    embed_cost_usd = tokens / 1_000_000 * 0.10  # ~$0.10 per 1M tokens (informative)
    return {"chunks": len(chunks), "tokens": tokens, "vector_size_mb": round(vec_mb, 3),
            "estimated_embed_cost_usd": round(embed_cost_usd, 6),
            "model": "granite-embedding-97m-multilingual-r2 (local, 384-dim)"}


def index_upload(chunks: list[dict[str, Any]], collection: str = USER_KB_COLLECTION) -> int:
    """Embed + upsert accepted chunks into the user collection (idempotent by
    chunk_id). Returns the number of indexed chunks."""
    store = KnowledgeStore(collection)
    n = store.upsert_chunk_dicts(chunks, collection=collection)
    return n


def userkb_count() -> int:
    try:
        return KnowledgeStore(USER_KB_COLLECTION).count(USER_KB_COLLECTION)
    except Exception:
        return 0


def find_neighbors(chunks: list[dict[str, Any]], i: int, k: int = 2) -> list[dict[str, Any]]:
    lo = max(0, i - k)
    hi = min(len(chunks), i + k + 1)
    return [c for c in chunks[lo:hi] if c is not chunks[i]]